import datetime as dt
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from ledger import Ledger, parse, correct_speech_entry

class LedgerFormatTests(unittest.TestCase):
    def setUp(self):
        self.temporary=tempfile.TemporaryDirectory()
        self.root=Path(self.temporary.name)/'桌面'/'账目'
        self.state=Path(self.temporary.name)/'状态'
        self.ledger=Ledger(self.root,self.state)
        self.when=dt.datetime(2026,10,13,12,34,56)

    def tearDown(self):
        self.ledger.close()
        self.temporary.cleanup()

    def old_book(self, extra=False, committed_id=None):
        self.root.mkdir(parents=True,exist_ok=True)
        book=Workbook()
        sheet=book.active
        sheet.title='账目'
        sheet.append(['系统时间','姓名','数量','类型']+(['自定义列'] if extra else []))
        sheet.append([self.when,'清风','3个','12A']+(['必须保留'] if extra else []))
        sheet.cell(2,1).number_format='yyyy-mm-dd hh:mm:ss'
        ids=book.create_sheet('_应用记录')
        ids.sheet_state='hidden'
        if committed_id:
            ids.append([committed_id])
        book.create_sheet('其他信息').append(['保留此页'])
        target=self.root/'清风.xlsx'
        book.save(target)
        book.close()
        return target,target.read_bytes()

    def test_new_three_column_workbook_and_append(self):
        self.ledger.enqueue(*parse('记账，清风，4个88A硒鼓'),when=self.when)
        self.ledger.enqueue('清风','2套12A墨盒',when=self.when)
        book=load_workbook(self.root/'清风.xlsx')
        sheet=book['账目']
        self.assertEqual(sheet.max_column,3)
        self.assertEqual(sheet.max_row,3)
        self.assertEqual([c.value for c in sheet[1]],['系统时间','姓名','账目'])
        self.assertEqual([c.value for c in sheet[2]],[self.when,'清风','4个88A硒鼓'])
        self.assertEqual(sheet['A2'].number_format,'yyyy-mm-dd hh:mm:ss')
        self.assertEqual(sheet['C3'].value,'2套12A墨盒')
        book.close()

    def test_old_workbook_merge_backup_and_other_sheet_preservation(self):
        target,before=self.old_book(committed_id='old-id')
        self.ledger.enqueue('清风','4个88A硒鼓',when=self.when)
        backups=list((self.root/'迁移备份').glob('*.xlsx'))
        self.assertEqual(len(backups),1)
        self.assertEqual(backups[0].read_bytes(),before)
        book=load_workbook(target)
        self.assertEqual(book['账目'].max_column,3)
        self.assertEqual(book['账目']['C2'].value,'3个12A')
        self.assertEqual(book['账目']['C3'].value,'4个88A硒鼓')
        self.assertEqual(book['其他信息']['A1'].value,'保留此页')
        self.assertEqual(book['_应用记录']['A1'].value,'old-id')
        book.close()
        self.ledger.enqueue('清风','1只88A硒鼓',when=self.when)
        self.assertEqual(len(list((self.root/'迁移备份').glob('*.xlsx'))),1)

    def test_old_sqlite_history_and_pending_recovery(self):
        self.ledger.close()
        state=Path(self.temporary.name)/'旧版状态'
        state.mkdir()
        db=sqlite3.connect(state/'待写入.sqlite3')
        db.execute('CREATE TABLE events (id TEXT PRIMARY KEY,stamp TEXT,name TEXT,qty TEXT,code TEXT,raw TEXT,done INTEGER DEFAULT 0,error TEXT DEFAULT "")')
        db.execute('INSERT INTO events VALUES(?,?,?,?,?,?,?,?)',('old-pending',str(self.when),'清风','2只','88A','旧记录',0,''))
        db.execute('INSERT INTO events VALUES(?,?,?,?,?,?,?,?)',('old-complete',str(self.when),'李凌宇','3个','12A','旧记录',1,''))
        db.commit();db.close()
        self.ledger=Ledger(self.root,state)
        self.assertEqual({r[0]:r[3] for r in self.ledger.recent()},{'old-pending':'2只88A','old-complete':'3个12A'})
        self.ledger.retry()
        self.assertTrue(all(r[4] for r in self.ledger.recent()))
        book=load_workbook(self.root/'清风.xlsx')
        self.assertEqual(book['账目']['C2'].value,'2只88A')
        book.close()

    def test_migration_retry_preserves_original_and_reuses_backup(self):
        target,before=self.old_book()
        with patch('ledger.os.replace',side_effect=PermissionError('locked')):
            uid=self.ledger.enqueue('清风','4个88A硒鼓',when=self.when)
        self.assertEqual(target.read_bytes(),before)
        self.assertEqual(self.ledger.db.execute('SELECT done FROM events WHERE id=?',(uid,)).fetchone()[0],0)
        self.ledger.retry()
        self.assertEqual(len(list((self.root/'迁移备份').glob('*.xlsx'))),1)
        self.ledger.db.execute('UPDATE events SET done=0 WHERE id=?',(uid,));self.ledger.db.commit()
        self.ledger.retry()
        book=load_workbook(target)
        self.assertEqual(book['账目'].max_row,3)
        self.assertEqual(book['账目'].max_column,3)
        book.close()

    def test_custom_columns_not_overwritten(self):
        target,before=self.old_book(extra=True)
        uid=self.ledger.enqueue('清风','4个88A硒鼓',when=self.when)
        self.assertEqual(target.read_bytes(),before)
        self.assertEqual(self.ledger.db.execute('SELECT done FROM events WHERE id=?',(uid,)).fetchone()[0],0)

    def test_text_that_starts_with_equals_is_not_formula(self):
        self.ledger.enqueue('清风','=1+1',when=self.when)
        book=load_workbook(self.root/'清风.xlsx')
        self.assertEqual(book['账目']['C2'].data_type,'s')
        self.assertEqual(book['账目']['C2'].value,'=1+1')
        book.close()

    def test_product_suffixes_chinese_numbers_and_long_descriptions(self):
        self.assertEqual(parse('记账，清风，4个88A硒鼓'),('清风','4个88A硒鼓'))
        self.assertEqual(parse('记账清风四个八十八A硒鼓。'),('清风','4个88A硒鼓'))
        self.assertEqual(parse('记账，清风，4个88A硒鼓，2只12A墨盒'),('清风','4个88A硒鼓,2只12A墨盒'))
        self.assertEqual(parse('记账清风四个88A硒鼓，质量正常'),('清风','4个88A硒鼓,质量正常'))
        self.assertEqual(parse('记账，清风，退货1个硒鼓'),('清风','退货1个硒鼓'))
        self.assertIsNone(parse('清风，4个88A硒鼓'))
        self.assertEqual(correct_speech_entry('4个88A硒谷'),'4个88A硒鼓')
        self.assertEqual(correct_speech_entry('4个88A硒鼓一只十二A硒鼓九个十二A碳粉'),'4个88A硒鼓1只12A硒鼓9个12A碳粉')
        self.assertEqual(correct_speech_entry('硒谷项目'),'硒谷项目')
        self.assertEqual(parse('记账，清风，4个88A硒谷'),('清风','4个88A硒谷'))

    def test_incomplete_and_invalid_names_rejected(self):
        for text in ['记账','记账，清风，','记账清风3个','记账，清风，0个88A','记账，../清风，4个88A硒鼓']:
            with self.assertRaises(ValueError,msg=text):
                parse(text)

if __name__=='__main__':
    unittest.main()
