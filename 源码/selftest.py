"""Tests run in disposable folders; never create sample entries on the real desktop."""
import datetime as dt
import json
import os
import socket
import tempfile
import wave
import io
import unittest
from pathlib import Path
from unittest.mock import patch
from ledger import Ledger, parse, number, correct_speech_entry

def run(model_path, audio_path=None):
    results=[]
    def check(label, condition):
        if not condition:
            raise AssertionError(label)
        results.append(label)
    def blocked(*args, **kwargs):
        raise AssertionError('测试期间禁止网络连接')
    with patch.object(socket.socket, 'connect', blocked), patch.object(socket, 'create_connection', blocked):
        from test_microphone_access import MicrophoneAccessTests
        from test_ledger_format import LedgerFormatTests
        from test_recording_window import RecordingWindowTests
        from test_capture_boundary import CaptureBoundaryTests
        from microphone_access import windows_access_status
        output=io.StringIO()
        suite=unittest.defaultTestLoader.loadTestsFromTestCase(MicrophoneAccessTests)
        mic_result=unittest.TextTestRunner(stream=output).run(suite)
        check('麦克风权限与设备检查的 8 项分支测试',mic_result.wasSuccessful())
        suite=unittest.defaultTestLoader.loadTestsFromTestCase(LedgerFormatTests)
        ledger_result=unittest.TextTestRunner(stream=output).run(suite)
        check('三列表格、旧版迁移与备份的 8 项测试',ledger_result.wasSuccessful())
        suite=unittest.defaultTestLoader.loadTestsFromTestCase(RecordingWindowTests)
        window_result=unittest.TextTestRunner(stream=output).run(suite)
        check('两分钟窗口、连续累积和计时边界的 5 项测试',window_result.wasSuccessful())
        suite=unittest.defaultTestLoader.loadTestsFromTestCase(CaptureBoundaryTests)
        capture_result=unittest.TextTestRunner(stream=output).run(suite)
        check('截止音频分割与最后一段语音整理的 2 项测试',capture_result.wasSuccessful())
        status=windows_access_status()
        check('Windows 麦克风访问状态接口可调用',status in (0,1,2,3))
        results.append('Windows 麦克风访问状态：'+{0:'未确定，使用实际设备检测',1:'允许',2:'用户拒绝',3:'系统拒绝'}[status])
        check('激活与硒鼓示例解析', parse('记账，清风，4个88A硒鼓') == ('清风','4个88A硒鼓'))
        check('没有激活文字不记账', parse('李凌宇，3个12A') is None)
        check('中文数字与字母发音', parse('请记账李凌宇三只八十八哎') == ('李凌宇','3只88A'))
        check('数量单位套', parse('记账张三十二套十二a') == ('张三','12套12A'))
        check('姓名纠正', parse('记账李玲宇3个12A',{'李玲宇':'李凌宇'})[0]=='李凌宇')
        check('中文数字', number('一万零三十') == 10030 and number('一二')==12)
        for bad in ['记账李凌宇0个12A','记账李凌宇3个','记账../李凌宇3个12A','记账，清风，']:
            try:
                parse(bad)
            except ValueError:
                pass
            else:
                raise AssertionError('应拒绝不完整或不安全输入：'+bad)
        results.append('拒绝不完整或不安全输入')
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder)/'桌面'/'账目'
            state=Path(folder)/'状态'
            ledger=Ledger(root,state)
            stamp=dt.datetime(2026,10,13,12,34,56)
            uid=ledger.enqueue('李凌宇','3个12A',when=stamp)
            ledger.enqueue('李凌宇','2只88A硒鼓',when=stamp)
            ledger.enqueue('张三','4套12A',when=stamp)
            wb=load_workbook(root/'李凌宇.xlsx')
            ws=wb['账目']
            check('新建文件夹与姓名工作簿', (root/'张三.xlsx').exists())
            check('追加两行与三列', ws.max_row==3 and ws.max_column==3)
            check('系统时间精确到秒',ws['A2'].value==stamp and ws['A2'].number_format=='yyyy-mm-dd hh:mm:ss')
            check('姓名和合并后的账目', [ws.cell(2,i).value for i in range(2,4)]==['李凌宇','3个12A'])
            wb.close()
            # Simulate crash after workbook commit but before journal confirmation.
            ledger.db.execute('UPDATE events SET done=0 WHERE id=?',(uid,)); ledger.db.commit()
            ledger.retry()
            wb=load_workbook(root/'李凌宇.xlsx')
            check('重试不会重复追加',wb['账目'].max_row==3)
            wb.close()
            real_replace=os.replace
            with patch('ledger.os.replace',side_effect=PermissionError('locked')):
                pending=ledger.enqueue('李凌宇','1个88A',when=stamp)
            check('占用时保留待写入',ledger.db.execute('SELECT done FROM events WHERE id=?',(pending,)).fetchone()[0]==0)
            ledger.close()
            ledger=Ledger(root,state)
            ledger.retry()
            check('重启后恢复写入',all(row[4] for row in ledger.recent()))
            wb=load_workbook(root/'李凌宇.xlsx')
            check('恢复后内容完整',wb['账目'].max_row==4)
            wb.close()
            ledger.close()
        import sounddevice as sd
        results.append(f'音频依赖加载成功，检测到 {len(sd.query_devices())} 个音频设备')
        import numpy as np
        from recognition import load_model, transcribe
        model,vad,window=load_model(model_path)
        transcribe(model,np.zeros(16000,dtype=np.float32))
        results.append('禁止网络连接情况下本地模型加载与推理成功')
        transcript=None
        if audio_path:
            with wave.open(str(audio_path),'rb') as audio:
                if audio.getnchannels()!=1 or audio.getsampwidth()!=2:
                    raise ValueError('测试音频必须为单声道 16 位 PCM')
                samples=np.frombuffer(audio.readframes(audio.getnframes()),dtype=np.int16).astype(np.float32)/32768
                transcript=transcribe(model,samples,audio.getframerate())
                aliases={'李玲雨':'李凌宇','李玲宇':'李凌宇','李凌雨':'李凌宇','李灵宇':'李凌宇'}
                record=parse(transcript,aliases)
                check('测试语音识别并解析为完整硒鼓账目',(record[0],correct_speech_entry(record[1]))==('清风','4个88A硒鼓'))
                rate=audio.getframerate()
                if rate!=16000:
                    samples=np.interp(np.arange(int(len(samples)*16000/rate))*rate/16000,np.arange(len(samples)),samples).astype(np.float32)
                samples=np.concatenate((samples,np.zeros(32000,dtype=np.float32)))
                for offset in range(0,len(samples)-window,window):
                    vad.accept_waveform(samples[offset:offset+window])
                check('语音端点检测产生完整语句',not vad.empty())
                detected=transcribe(model,np.asarray(vad.front.samples,dtype=np.float32))
                record=parse(detected,aliases)
                check('麦克风同款分段链路可解析硒鼓示例',(record[0],correct_speech_entry(record[1]))==('清风','4个88A硒鼓'))
        return {'passed':True,'checks':results,'speech_transcript':transcript}

if __name__=='__main__':
    import sys
    base=Path(__file__).resolve().parent.parent
    result=run(base/'模型'/'sensevoice',sys.argv[1] if len(sys.argv)>1 else None)
    print(json.dumps(result,ensure_ascii=False,indent=2))
