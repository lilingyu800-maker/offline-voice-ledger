"""Offline ledger domain logic. No network and no microphone dependencies."""
import ctypes
import datetime as dt
import json
import os
import re
import sqlite3
import tempfile
import unicodedata
import uuid
import hashlib
from pathlib import Path

DIGITS = dict(zip('零〇一二两三四五六七八九', [0,0,1,2,2,3,4,5,6,7,8,9]))
NUMBER = r'[0-9零〇一二两三四五六七八九十百千万]+'

def number(text):
    if text.isdigit():
        return int(text)
    if all(c in DIGITS for c in text):
        return int(''.join(str(DIGITS[c]) for c in text))
    result = section = digit = 0
    for c in text:
        if c in DIGITS:
            digit = DIGITS[c]
        elif c == '万':
            result += (section + digit) * 10000
            section = digit = 0
        elif c in '十百千':
            section += (digit or 1) * {'十':10,'百':100,'千':1000}[c]
            digit = 0
        else:
            raise ValueError('无法识别数字')
    return result + section + digit

def compact(text):
    return re.sub(r'[\s，,。；;：:、！!？?]', '', unicodedata.normalize('NFKC', text))

def parse(text, aliases=None, require_wake=True):
    text = unicodedata.normalize('NFKC',text).strip()
    if require_wake:
        if '记账' not in text:
            return None
        text = text.split('记账', 1)[1]
    text=text.strip(' \t,，:：;；、。！!？?')
    # Prefer an explicit name/content separator; speech often has no punctuation,
    # in which case a configured name or the first quantity+unit gives the boundary.
    parts=re.split(r'[,，:：;；、]',text,maxsplit=1)
    if len(parts)==2 and compact(parts[0]) and not re.search(NUMBER+r'(?:个|只|套)',parts[0]):
        name,entry=compact(parts[0]),parts[1].strip()
    else:
        clean=re.sub(r'\s+','',text)
        known=sorted((key for key in (aliases or {}) if clean.startswith(key)),key=len,reverse=True)
        if known:
            name,entry=known[0],clean[len(known[0]):]
        else:
            m=re.fullmatch(r'(?P<name>.{2,}?)(?P<entry>'+NUMBER+r'(?:个|只|套).+)',clean)
            if not m:
                raise ValueError('请说：记账，姓名，账目内容，例如“记账，清风，4个88A硒鼓”')
            name,entry=m['name'],m['entry']
    name = (aliases or {}).get(name, name)
    validate_name(name)
    entry=normalize_entry(entry)
    return name,entry

def validate_entry(entry):
    if not entry or not entry.strip() or len(entry)>32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',entry):
        raise ValueError('账目内容不能为空、含有非法控制字符或超过 Excel 单元格长度上限')

def normalize_entry(entry):
    entry=unicodedata.normalize('NFKC',entry).strip().strip('，,。；;！!？?')
    entry=re.sub(r'\s+','',entry)
    # Keep the entire description, including products such as 硒鼓. Only normalize
    # leading spoken quantities and model numbers as in previous versions.
    qty=re.match(r'('+NUMBER+r')(个|只|套)',entry)
    if qty:
        value=number(qty[1])
        if not 1<=value<=999999:
            raise ValueError('数量应在 1 到 999999 之间')
        rest=entry[qty.end():]
        if not rest:
            raise ValueError('请补充数量后面的账目内容')
        rest=re.sub(r'^('+NUMBER+r')(?:诶伊|诶呀|诶|哎|欸|唉)(?=$|硒鼓|墨盒|粉盒)',lambda m:f'{number(m[1])}A',rest)
        rest=re.sub(r'^('+NUMBER+r')([a-zA-Z])',lambda m:f'{number(m[1])}{m[2].upper()}',rest)
        entry=f'{value}{qty[2]}{rest}'
    validate_entry(entry)
    return entry

def correct_speech_entry(entry):
    entry=re.sub(r'('+NUMBER+r')(个|只|套)('+NUMBER+r')([a-zA-Z])',lambda m:f'{number(m[1])}{m[2]}{number(m[3])}{m[4].upper()}',entry)
    # Chinese ASR can choose the homophone 硒谷 for the product 硒鼓.
    # Apply this narrowly to a model suffix in voice input; manual text is literal.
    return re.sub(r'(?<=[0-9A-Za-z])硒谷(?=$|[,，。；;])','硒鼓',entry)

def validate_name(name):
    if not name or len(name) > 80 or re.search(r'[<>:"/\\|?*\x00-\x1f]', name) or name.endswith(('.', ' ')):
        raise ValueError('姓名为空、过长或含有 Windows 文件名不允许的字符')
    if name.split('.')[0].upper() in {'CON','PRN','AUX','NUL',*(f'COM{i}' for i in range(1,10)),*(f'LPT{i}' for i in range(1,10))}:
        raise ValueError('姓名不能使用 Windows 保留的文件名')

def desktop():
    if os.name == 'nt':
        buf = ctypes.create_unicode_buffer(32768)
        if ctypes.windll.shell32.SHGetFolderPathW(None, 0x10, None, 0, buf) == 0:
            return Path(buf.value)
    return Path.home() / 'Desktop'

def state_dir():
    folder = Path(os.environ.get('LOCALAPPDATA', Path.home())) / 'OfflineVoiceLedger'
    folder.mkdir(parents=True, exist_ok=True)
    return folder

class Ledger:
    def __init__(self, root=None, state=None):
        self.root = Path(root) if root else desktop() / '账目'
        self.state = Path(state) if state else state_dir()
        self.state.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(self.state / '待写入.sqlite3')
        self.db.execute('CREATE TABLE IF NOT EXISTS events (id TEXT PRIMARY KEY, stamp TEXT, name TEXT, qty TEXT, code TEXT, raw TEXT, done INTEGER DEFAULT 0, error TEXT DEFAULT "")')
        columns={row[1] for row in self.db.execute('PRAGMA table_info(events)')}
        if 'content' not in columns:
            self.db.execute('ALTER TABLE events ADD COLUMN content TEXT')
        self.db.execute('UPDATE events SET content=COALESCE(qty,"")||COALESCE(code,"") WHERE content IS NULL')
        self.db.commit()

    def enqueue(self, name, entry, raw='', when=None, uid=None):
        validate_name(name)
        validate_entry(entry)
        uid = uid or uuid.uuid4().hex
        existing=self.db.execute('SELECT name,content FROM events WHERE id=?',(uid,)).fetchone()
        if existing:
            if existing!=(name,entry):
                raise ValueError('这条记录已保留，请使用“重试未写入”，或在工作簿中修正已写入内容')
            self.write(uid)
            return uid
        stamp = (when or dt.datetime.now()).replace(microsecond=0).isoformat(sep=' ')
        self.db.execute('INSERT INTO events(id,stamp,name,qty,code,content,raw) VALUES(?,?,?,?,?,?,?)', (uid,stamp,name,'','',entry,raw))
        self.db.commit()
        self.write(uid)
        return uid

    def write(self, uid):
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        row = self.db.execute('SELECT id,stamp,name,content,done FROM events WHERE id=?', (uid,)).fetchone()
        if not row or row[4]:
            return
        book = None
        tmp = None
        try:
            self.root.mkdir(parents=True, exist_ok=True)
            target = self.root / f'{row[2]}.xlsx'
            migrated=False
            if target.exists():
                book = load_workbook(target)
                visible = [s for s in book.worksheets if s.sheet_state == 'visible']
                sheet = book['账目'] if '账目' in book.sheetnames else visible[0]
                headers=list(next(sheet.iter_rows(min_row=1,max_row=1,values_only=True)))
                if headers==['系统时间','姓名','数量','类型']:
                    if sheet.max_column>4:
                        raise ValueError('旧表格含有额外列，请先核对格式，避免迁移影响自定义内容')
                    # Preserve the original bytes before combining the old columns.
                    original=target.read_bytes()
                    backups=self.root/'迁移备份'
                    backups.mkdir(exist_ok=True)
                    backup=backups/f'{target.stem}_四列备份_{hashlib.sha256(original).hexdigest()[:12]}.xlsx'
                    if not backup.exists():
                        with backup.open('xb') as stream:
                            stream.write(original)
                    for index in range(2,sheet.max_row+1):
                        qty,code=sheet.cell(index,3).value,sheet.cell(index,4).value
                        sheet.cell(index,3).value=('' if qty is None else str(qty))+('' if code is None else str(code))
                        sheet.cell(index,3).data_type='s'
                    sheet.delete_cols(4)
                    sheet.column_dimensions.pop('D',None)
                    sheet.cell(1,3).value='账目'
                    if sheet.auto_filter.ref:
                        sheet.auto_filter.ref=f'A1:C{sheet.max_row}'
                    migrated=True
                elif headers!=['系统时间','姓名','账目']:
                    raise ValueError('现有表格表头应为“系统时间、姓名、账目”；记录已保留，请先核对表格格式')
            else:
                book = Workbook()
                sheet = book.active
                sheet.title = '账目'
                sheet.append(['系统时间', '姓名', '账目'])
                sheet.freeze_panes = 'A2'
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='2563EB')
                    cell.alignment = Alignment(vertical='center')
                sheet.row_dimensions[1].height = 26
                for col, width in [('A',24),('B',18),('C',60)]:
                    sheet.column_dimensions[col].width = width
            sheet.column_dimensions['C'].width=60
            ids = book['_应用记录'] if '_应用记录' in book.sheetnames else book.create_sheet('_应用记录')
            ids.sheet_state = 'hidden'
            if not any(cell[0] == uid for cell in ids.iter_rows(values_only=True)):
                sheet.append([dt.datetime.fromisoformat(row[1]),row[2],row[3]])
                sheet.cell(sheet.max_row, 1).number_format = 'yyyy-mm-dd hh:mm:ss'
                # Force literal strings, so names starting with '=' are not formulas.
                for col in (2,3):
                    sheet.cell(sheet.max_row, col).data_type = 's'
                sheet.cell(sheet.max_row,3).alignment=Alignment(wrap_text=True,vertical='top')
                ids.append([uid])
                migrated=True
            if migrated:
                fd, tmp = tempfile.mkstemp(prefix='.账目-', suffix='.xlsx', dir=self.root)
                os.close(fd)
                book.save(tmp)
                book.close()
                book = None
                os.replace(tmp, target)
                tmp = None
            self.db.execute('UPDATE events SET done=1,error="" WHERE id=?', (uid,))
        except Exception as exc:
            msg = '文件被 Excel/WPS 占用，请关闭该文件后点击“重试未写入”' if isinstance(exc, PermissionError) else str(exc)
            self.db.execute('UPDATE events SET error=? WHERE id=?', (msg,uid))
        finally:
            if book:
                book.close()
            if tmp and Path(tmp).exists():
                Path(tmp).unlink()
            self.db.commit()

    def retry(self):
        for (uid,) in self.db.execute('SELECT id FROM events WHERE done=0 ORDER BY rowid').fetchall():
            self.write(uid)

    def recent(self):
        return self.db.execute('SELECT id,stamp,name,content,done,error FROM events ORDER BY rowid DESC LIMIT 100').fetchall()

    def close(self):
        self.db.close()
